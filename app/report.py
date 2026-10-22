# -*- coding: utf-8 -*-
"""导出 Word 版审查意见书。

产出物是「给业务方看」的文档，不是界面截图的堆砌，所以按正式意见书结构写：
结论 -> 统计 -> 逐条明细（依据 / 命中原文 / 风险说明 / 修改建议）-> 免责声明。
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .checklist import REMEDY_LABEL

SEVERITY_COLOR = {
    "high": RGBColor(0xC0, 0x2A, 0x2A),
    "medium": RGBColor(0xB5, 0x6A, 0x00),
    "low": RGBColor(0x4A, 0x6A, 0x8A),
}

# 签署前必检清单：借鉴 contract-copilot 的 review-framework。
# (项目, 关联检查点, 无法自动判断时的说明)
PRE_SIGN_CHECKS = [
    ("核心条款（价款 / 期限 / 生效）齐全", ["CP11", "CP12", "CP18"], ""),
    ("付款、交付、验收形成闭环", ["CP08", "CP22", "CP25"], ""),
    ("违约与解除机制可直接执行", ["CP02", "CP03", "CP04", "CP14", "CP21"], ""),
    ("争议解决条款完整、机构与管辖明确", ["CP06", "CP07", "CP33"], ""),
    ("通知、送达与证据机制已约定", ["CP17"], ""),
    ("保密范围与期限清晰", ["CP15"], ""),
    ("主体资格与授权文件已核验", [], "需人工核验"),
    ("标的合法性与可处分状态已核验", [], "需人工核验"),
]


def _set_base_font(document):
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    try:
        rpr = style.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001 - 字体缺失不影响文档有效性
        pass


def _para(document, text="", *, size=10.5, bold=False, color=None, align=None,
          space_before=0, space_after=4):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.35
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = "宋体"
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001
        pass
    return paragraph


def _heading(document, text):
    return _para(document, text, size=13, bold=True, space_before=12, space_after=6)


def build(result: dict) -> bytes:
    document = Document()
    _set_base_font(document)

    _para(document, "合同审查意见书", size=20, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(document, "Contract Review Report", size=9,
          align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x88, 0x88, 0x88), space_after=14)

    meta = [
        ("审查编号", result.get("review_id", "-")),
        ("文件名称", result.get("filename", "-")),
        ("生成时间", result.get("created_at", "-")),
        ("文本规模", f'{result.get("contract", {}).get("chars", 0)} 字 / '
                     f'{result.get("contract", {}).get("paragraphs", 0)} 段'),
        ("检索引擎", "规则层" + (
            f'＋大模型（{result["engine"].get("llm_model", "")}）'
            if result.get("engine", {}).get("llm") else "")),
    ]
    identity = document.add_table(rows=0, cols=2)
    identity.style = "Table Grid"
    for key, value in meta:
        row = identity.add_row().cells
        row[0].text = key
        row[1].text = str(value)
        for cell, bold in ((row[0], True), (row[1], False)):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.bold = bold

    _heading(document, "一、总体结论")
    score = result.get("score", 0)
    level = result.get("risk_level", "-")
    color = SEVERITY_COLOR["high"] if score < 50 else (
        SEVERITY_COLOR["medium"] if score < 70 else SEVERITY_COLOR["low"])
    _para(document, f"综合评分：{score} / 100　　风险等级：{level}", size=12, bold=True, color=color)
    _para(document, result.get("advice", ""))
    _para(document, result.get("overall_comment", ""))

    stats = result.get("stats", {})
    _para(document, f'风险分布：高风险 {stats.get("high", 0)} 项，'
                    f'中风险 {stats.get("medium", 0)} 项，'
                    f'低风险 {stats.get("low", 0)} 项，合计 {stats.get("total", 0)} 项。',
          space_before=4)

    _heading(document, "二、风险明细")
    findings = result.get("findings") or []
    if not findings:
        _para(document, "未命中预设审查点。请注意这不代表合同无风险，仍建议人工复核。")

    for index, finding in enumerate(findings, 1):
        severity = finding.get("severity", "low")
        _para(document, f'{index}. {finding.get("name", "")}'
                        f'（{finding.get("severity_label", "")}风险 · {finding.get("category", "")}）',
              size=11.5, bold=True, color=SEVERITY_COLOR.get(severity), space_before=10, space_after=3)

        remedy = finding.get("remedy") or ""
        if remedy:
            _para(document, f"整改动作：{REMEDY_LABEL.get(remedy, remedy)}", size=9.5,
                  color=RGBColor(0x55, 0x55, 0x55))

        statutes = finding.get("statutes") or []
        if statutes:
            cites = "；".join(f'《民法典》{s["article"]}' for s in statutes)
            _para(document, f"法律依据：{cites}", size=9.5)
        else:
            _para(document, "法律依据：本条为通用条款设计建议，无直接对应条文。", size=9.5)

        evidence = finding.get("evidence") or []
        if evidence:
            quoted = evidence[0]["text"]
            if len(quoted) > 160:
                quoted = quoted[:160] + "……"
            _para(document, f"命中原文：「{quoted}」", size=9.5, color=RGBColor(0x55, 0x55, 0x55))
        else:
            _para(document, "命中方式：属于条款缺失或约定不明，未标注原文位置。", size=9.5,
                  color=RGBColor(0x55, 0x55, 0x55))

        _para(document, f"风险说明：{finding.get('rationale', '')}")

        verdict = finding.get("llm_verdict")
        if verdict == "confirm":
            confidence = finding.get("llm_confidence") or 0
            _para(document, f"复核结论：确认构成风险（置信度 {confidence:.0%}）。", size=9.5)
        elif verdict == "dismiss":
            _para(document, "复核结论：大模型复核认为不构成实质风险，建议人工确认后忽略本条。",
                  size=9.5, color=RGBColor(0x77, 0x77, 0x77))
        elif verdict == "uncertain":
            _para(document, "复核结论：信息不足以判断，需补充材料或人工确认。", size=9.5,
                  color=SEVERITY_COLOR["medium"])

        if finding.get("llm_analysis"):
            _para(document, f"结合本合同：{finding['llm_analysis']}")
        _para(document, f"修改建议：{finding.get('suggestion', '')}")
        if finding.get("revised_clause"):
            _para(document, "建议条款文本：", bold=True, space_before=2)
            _para(document, finding["revised_clause"], color=RGBColor(0x22, 0x44, 0x66))

    _heading(document, "三、谈判顺序")
    if findings:
        _para(document, "按风险后果排序处理，先解决会导致条款无效或无法履行的问题，再谈优化表达。",
              size=9.5, color=RGBColor(0x55, 0x55, 0x55))
        order = {"high": 0, "medium": 1, "low": 2}
        ranked = sorted(findings, key=lambda f: (order.get(f.get("severity"), 9), f.get("id", "")))
        for rank, finding in enumerate(ranked, 1):
            remedy = REMEDY_LABEL.get(finding.get("remedy") or "", "待定")
            _para(document,
                  f'{rank}. {finding.get("name", "")}'
                  f'（{finding.get("severity_label", "")}风险，建议{remedy}）',
                  size=9.5, space_after=2)
    else:
        _para(document, "无需分级谈判，按常规复核流程签署即可。", size=9.5)

    _heading(document, "四、签署前必检清单")
    flagged = {f.get("id") for f in findings}
    for label, related, manual in PRE_SIGN_CHECKS:
        if manual:
            mark, hint = "○", manual
        else:
            hits = [rid for rid in related if rid in flagged]
            if hits:
                mark, hint = "✗", "未通过：" + "、".join(hits)
            else:
                mark, hint = "✓", "已覆盖"
        _para(document, f"{mark} {label}　{hint}", size=9.5, space_after=2)

    _heading(document, "五、免责声明")
    _para(document, result.get("disclaimer", ""), size=9, color=RGBColor(0x77, 0x77, 0x77))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
