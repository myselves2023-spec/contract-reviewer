# -*- coding: utf-8 -*-
"""文档解析：把 .docx / .txt 转成带字符偏移的分段结构。

设计要点：不还原版式，只保证「文本 + 偏移」。偏移是后续条款级高亮定位的基础，
所以这里的 start/end 必须与返回的 text 严格一致。
"""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

# PyMuPDF 装在项目内的 vendor/ 目录里，避免污染系统的 Python 环境。
# 没有 vendor/ 时 PDF 分支给出明确提示，其余格式不受影响。
_VENDOR = Path(__file__).resolve().parent.parent / "vendor"
if _VENDOR.is_dir() and str(_VENDOR) not in sys.path:
    sys.path.append(str(_VENDOR))


class ParseError(Exception):
    """解析失败，message 会直接展示给用户。"""


SUPPORTED_EXT = (".docx", ".txt", ".md", ".markdown", ".pdf")
MIN_PDF_TEXT = 50  # 少于此字符数视为无文字层（扫描件）


def _clean(paragraphs):
    out = []
    for p in paragraphs:
        t = (p or "").replace("\ufeff", "").replace("\u00a0", " ").replace("\u3000", " ").strip()
        if t:
            out.append(t)
    return out


def _from_docx(data: bytes):
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ParseError(
            "服务端缺少 python-docx，无法解析 .docx。请改用 .txt 上传，"
            "或在部署环境执行 pip install python-docx。"
        ) from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("该 .docx 无法打开，可能已损坏或不是 Word 文档。") from exc

    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and (not cells or cells[-1] != value):
                    cells.append(value)
            if cells:
                paragraphs.append(" | ".join(cells))
    return paragraphs


def _from_text(data: bytes):
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding).splitlines()
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文件编码，请另存为 UTF-8 后重试。")


def _from_pdf(data: bytes):
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ParseError(
            "未安装 PDF 解析库。在项目根目录执行：\n"
            "  python -m pip install --target vendor pymupdf\n"
            "或改用 .docx / .txt 上传。"
        ) from exc

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ParseError("无法打开该 PDF，可能已损坏或不是有效的 PDF 文件。") from exc

    if document.needs_pass:
        if not document.authenticate(""):
            document.close()
            raise ParseError("该 PDF 已加密，请先解除密码保护后再上传。")

    paragraphs = []
    try:
        for page in document:
            text = page.get_text("text") or ""
            paragraphs.extend(text.splitlines())
    finally:
        document.close()

    total = sum(len(p.strip()) for p in paragraphs)
    if total < MIN_PDF_TEXT:
        raise ParseError(
            "这份 PDF 里提取不到文字，可能是扫描件或纯图片。"
            "本工具不做 OCR，请提供带文字层的 PDF，或先转成 .docx。"
        )
    return paragraphs


def extract(filename: str, data: bytes) -> dict:
    """返回 {filename, text, segments[], chars, paragraphs}。"""
    ext = os.path.splitext(filename or "")[1].lower()

    if ext == ".docx":
        paragraphs = _from_docx(data)
    elif ext in (".txt", ".md", ".markdown"):
        paragraphs = _from_text(data)
    elif ext == ".pdf":
        paragraphs = _from_pdf(data)
    elif ext == ".doc":
        raise ParseError("不支持旧版 .doc。请用 Word 另存为 .docx 后重新上传。")
    else:
        raise ParseError(f"不支持的文件类型「{ext or '未知'}」，请上传 .docx / .pdf / .txt / .md。")

    paragraphs = _clean(paragraphs)
    if not paragraphs:
        raise ParseError("未从文件中提取到正文文本（可能是扫描件或纯图片文档）。")

    segments = []
    parts = []
    pos = 0
    for paragraph in paragraphs:
        start = pos
        parts.append(paragraph)
        pos += len(paragraph)
        segments.append({"index": len(segments), "text": paragraph, "start": start, "end": pos})
        parts.append("\n")
        pos += 1

    return {
        "filename": filename,
        "text": "".join(parts),
        "segments": segments,
        "chars": pos,
        "paragraphs": len(segments),
    }
