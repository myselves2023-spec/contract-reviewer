# -*- coding: utf-8 -*-
"""把 samples/*.txt 批量转成排版像样的 Word，用于演示上传流程。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"

# 文件名映射：源 txt -> 导出 docx
TARGETS = {
    "sample_contract.txt": "技术服务合同_样例.docx",
    "租赁合同_样例.txt": "租赁合同_样例.docx",
}

CLAUSE_PREFIXES = ("第",)
SIGN_PREFIXES = ("甲方：", "乙方：", "丙方：", "出租方：", "承租方：")


def style_run(run, size=12, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "宋体"
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001
        pass


def convert(src: Path, dst: Path) -> int:
    lines = [line.strip() for line in src.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    document = Document()

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(line)

        is_clause = line.startswith(CLAUSE_PREFIXES) and "条" in line[:6] and len(line) < 30
        is_sign = line.startswith(SIGN_PREFIXES)

        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(run, size=18, bold=True)
            paragraph.paragraph_format.space_after = Pt(16)
        elif is_sign:
            style_run(run, size=12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif is_clause:
            style_run(run, size=12.5, bold=True)
            paragraph.paragraph_format.space_before = Pt(10)
        else:
            style_run(run, size=12)

    document.save(dst)
    return len(lines)


def main() -> int:
    count = 0
    for name, output in TARGETS.items():
        src = SAMPLES / name
        if not src.exists():
            print(f"  跳过（不存在）：{name}")
            continue
        dst = SAMPLES / output
        paragraphs = convert(src, dst)
        print(f"  {name} -> {output}  ({paragraphs} 段, {dst.stat().st_size} bytes)")
        count += 1
    print(f"完成 {count} 份")
    return 0


if __name__ == "__main__":
    sys.exit(main())
